"""Importacion y exportacion de inventario en Excel.

Pensado para trabajar sin conexion: se exporta el inventario, se edita en el
computador o en el celular, y se vuelve a subir.

La importacion ocurre en dos fases. Primero se analiza el archivo y se devuelve
lo que pasaria, sin tocar la base. Solo si el usuario confirma se aplica. Asi
nadie sobrescribe su inventario a ciegas.

Las anomalias se senalan, nunca se corrigen solas: un precio de venta por debajo
del costo o un salto brusco en el costo se informan, pero el dato entra tal como
lo escribio el usuario, que es quien sabe si es un error o una decision.
"""
import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .models.insumo import Insumo
from .models.precio_insumo import PrecioInsumo
from .models.proveedor import Proveedor

HOJA = "Inventario"

COLUMNAS = [
    ("nombre", "Producto", 34),
    ("categoria", "Categoría", 16),
    ("presentacion", "Presentación", 20),
    ("stock_actual", "Stock actual", 13),
    ("stock_minimo", "Stock mínimo", 13),
    ("costo", "Costo unitario (COP)", 19),
    ("precio_venta", "Precio venta (COP)", 17),
    ("proveedor", "Proveedor", 24),
]

# Variacion de costo a partir de la cual se pide confirmar que no es un error.
SALTO_COSTO = Decimal("0.5")

# Titulos alternativos aceptados al leer, para no romper archivos ya descargados.
ALIAS = {
    "costo": {"costo unitario (cop)", "costo (cop)", "costo unitario", "costo"},
    "precio_venta": {"precio venta (cop)", "precio de venta (cop)", "precio venta", "precio cliente"},
    "nombre": {"producto", "nombre", "insumo"},
}


def _texto_presentacion(insumo, mejor) -> str:
    """Describe la presentacion aclarando cuantas unidades trae."""
    if mejor is None:
        return insumo.unidad_medida or ""
    base = mejor.descripcion_presentacion or insumo.unidad_medida or ""
    unidades = mejor.unidades_por_presentacion or 1
    if unidades > 1:
        detalle = f"{unidades} unidades"
        return f"{base} ({detalle})" if base else detalle
    return base


# ---------------------------------------------------------------- exportar

def exportar_inventario(db: Session, espacio_id, nombre_espacio: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA

    encabezado = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="2563EB")

    ws.append([titulo for _, titulo, _ in COLUMNAS])
    for i, (_, _, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = encabezado
        celda.fill = fondo
        celda.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho

    insumos = (
        db.query(Insumo)
        .filter(Insumo.espacio_id == espacio_id, Insumo.activo == True)
        .order_by(Insumo.nombre)
        .all()
    )

    for insumo in insumos:
        activos = [p for p in insumo.precios if p.activo]
        mejor = min(activos, key=lambda p: p.precio_unitario) if activos else None
        ws.append([
            insumo.nombre,
            insumo.categoria or "",
            _texto_presentacion(insumo, mejor),
            float(insumo.stock_actual or 0),
            float(insumo.stock_minimo or 0),
            round(float(mejor.precio_unitario), 2) if mejor else None,
            float(insumo.precio_venta) if insumo.precio_venta is not None else None,
            mejor.proveedor.nombre if mejor and mejor.proveedor else "",
        ])

    ws.freeze_panes = "A2"

    ayuda = wb.create_sheet("Instrucciones")
    for fila in [
        ["Cómo usar este archivo"],
        [],
        ["1.", "Edita la hoja 'Inventario' en tu computador, sin necesidad de internet."],
        ["2.", "Para cambiar algo, escribe el valor nuevo sobre el anterior."],
        ["3.", "Para agregar un producto, escríbelo en una fila nueva al final."],
        ["4.", "No cambies los títulos de las columnas ni el nombre de la hoja."],
        ["5.", "Sube el archivo en la sección Inventario de la aplicación."],
        [],
        ["Antes de guardar nada, la aplicación te muestra qué va a cambiar."],
        ["Los productos que borres del archivo NO se borran de la aplicación."],
        [],
        ["El costo es por UNIDAD, no por caja. Si una caja de 100 guantes"],
        ["", "cuesta 45.200, el costo unitario es 452."],
        [],
        ["Los precios se escriben en pesos, sin puntos ni símbolo: 147000"],
        [f"Exportado desde: {nombre_espacio}"],
    ]:
        ayuda.append(fila)
    ayuda.column_dimensions["A"].width = 5
    ayuda.column_dimensions["B"].width = 70
    ayuda["A1"].font = Font(bold=True, size=13)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


# ---------------------------------------------------------------- lectura

def _a_decimal(valor: Any) -> Decimal | None:
    """Acepta 147000, '147.000', '$ 147.000,50' y variantes."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))

    texto = str(valor).strip()
    texto = re.sub(r"[^\d,.\-]", "", texto)
    if not texto:
        return None

    # Con ambos separadores, el ultimo que aparece es el decimal.
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        # Una sola coma con dos decimales es separador decimal; si no, de miles.
        texto = texto.replace(",", "." if len(texto.split(",")[-1]) == 2 else "")
    elif texto.count(".") > 1:
        texto = texto.replace(".", "")
    elif "." in texto and len(texto.split(".")[-1]) == 3:
        # 147.000 en Colombia son ciento cuarenta y siete mil, no 147 con decimales.
        texto = texto.replace(".", "")

    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def leer_filas(contenido: bytes) -> list[dict]:
    """Extrae las filas de la hoja de inventario, tolerando el orden de columnas."""
    try:
        wb = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as exc:
        raise ValueError(f"No se pudo leer el archivo: {exc}") from exc

    ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[0]

    titulos = {}
    for indice, celda in enumerate(ws[1]):
        if celda.value is None:
            continue
        texto = str(celda.value).strip().lower()
        for clave, titulo, _ in COLUMNAS:
            aceptados = ALIAS.get(clave, {titulo.strip().lower()})
            if texto == titulo.strip().lower() or texto in aceptados:
                titulos[clave] = indice
                break

    if "nombre" not in titulos:
        raise ValueError(
            "El archivo no tiene la columna 'Producto'. "
            "Descarga la plantilla desde la aplicación y trabaja sobre ella."
        )

    filas = []
    for numero, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in fila):
            continue

        def valor(clave):
            i = titulos.get(clave)
            return fila[i] if i is not None and i < len(fila) else None

        nombre = valor("nombre")
        filas.append({
            "fila": numero,
            "nombre": str(nombre).strip() if nombre is not None else "",
            "categoria": (str(valor("categoria")).strip() if valor("categoria") else None),
            "presentacion": (str(valor("presentacion")).strip() if valor("presentacion") else None),
            "stock_actual": _a_decimal(valor("stock_actual")),
            "stock_minimo": _a_decimal(valor("stock_minimo")),
            "costo": _a_decimal(valor("costo")),
            "precio_venta": _a_decimal(valor("precio_venta")),
            "proveedor": (str(valor("proveedor")).strip() if valor("proveedor") else None),
        })
    return filas


# ---------------------------------------------------------------- analisis

def _costo_actual(insumo: Insumo) -> Decimal | None:
    activos = [p for p in insumo.precios if p.activo]
    if not activos:
        return None
    return Decimal(str(min(p.precio_unitario for p in activos)))


def analizar(db: Session, espacio_id, contenido: bytes) -> dict:
    """Describe lo que haria la importacion, sin tocar la base de datos."""
    filas = leer_filas(contenido)

    existentes = {
        i.nombre.strip().lower(): i
        for i in db.query(Insumo).filter(
            Insumo.espacio_id == espacio_id, Insumo.activo == True
        ).all()
    }

    vistos: dict[str, int] = {}
    resultado = []

    for fila in filas:
        nombre = fila["nombre"]
        clave = nombre.lower()
        alertas: list[dict] = []
        cambios: list[str] = []

        if not nombre:
            resultado.append({
                "fila": fila["fila"], "nombre": "", "accion": "ignorar",
                "cambios": [], "alertas": [{
                    "severidad": "alta",
                    "mensaje": "La fila no tiene nombre de producto y se va a omitir",
                }],
            })
            continue

        if clave in vistos:
            alertas.append({
                "severidad": "alta",
                "mensaje": f"'{nombre}' ya aparece en la fila {vistos[clave]}. Se usará el último valor.",
            })
        vistos[clave] = fila["fila"]

        if fila["stock_actual"] is not None and fila["stock_actual"] < 0:
            alertas.append({"severidad": "alta", "mensaje": "El stock es negativo"})

        if fila["costo"] is not None and fila["precio_venta"] is not None:
            if fila["precio_venta"] < fila["costo"]:
                perdida = fila["costo"] - fila["precio_venta"]
                alertas.append({
                    "severidad": "alta",
                    "mensaje": f"El precio de venta está ${perdida:,.0f} por debajo del costo",
                })

        insumo = existentes.get(clave)

        if insumo is None:
            accion = "crear"
            if fila["costo"] is None:
                alertas.append({"severidad": "media", "mensaje": "Producto nuevo sin costo"})
        else:
            accion = "actualizar"

            stock_previo = Decimal(str(insumo.stock_actual or 0))
            if fila["stock_actual"] is not None and fila["stock_actual"] != stock_previo:
                cambios.append(f"stock {stock_previo:g} → {fila['stock_actual']:g}")

            minimo_previo = Decimal(str(insumo.stock_minimo or 0))
            if fila["stock_minimo"] is not None and fila["stock_minimo"] != minimo_previo:
                cambios.append(f"mínimo {minimo_previo:g} → {fila['stock_minimo']:g}")

            costo_previo = _costo_actual(insumo)
            if fila["costo"] is not None and (costo_previo is None or fila["costo"] != costo_previo):
                if costo_previo is None:
                    cambios.append(f"costo nuevo ${fila['costo']:,.0f}")
                else:
                    cambios.append(f"costo ${costo_previo:,.0f} → ${fila['costo']:,.0f}")
                    if costo_previo > 0:
                        variacion = abs(fila["costo"] - costo_previo) / costo_previo
                        if variacion > SALTO_COSTO:
                            direccion = "subió" if fila["costo"] > costo_previo else "bajó"
                            alertas.append({
                                "severidad": "media",
                                "mensaje": f"El costo {direccion} {variacion * 100:.0f}%. Verifica que no sea un error.",
                            })

            venta_previa = Decimal(str(insumo.precio_venta)) if insumo.precio_venta is not None else None
            if fila["precio_venta"] is not None and fila["precio_venta"] != venta_previa:
                if venta_previa is None:
                    cambios.append(f"precio de venta ${fila['precio_venta']:,.0f}")
                else:
                    cambios.append(f"precio de venta ${venta_previa:,.0f} → ${fila['precio_venta']:,.0f}")

            if fila["categoria"] and fila["categoria"] != (insumo.categoria or ""):
                cambios.append(f"categoría → {fila['categoria']}")

            if not cambios:
                accion = "sin_cambios"

        resultado.append({
            "fila": fila["fila"],
            "nombre": nombre,
            "accion": accion,
            "cambios": cambios,
            "alertas": alertas,
        })

    en_archivo = set(vistos)
    ausentes = sorted(i.nombre for clave, i in existentes.items() if clave not in en_archivo)

    return {
        "filas": resultado,
        "resumen": {
            "leidas": len(filas),
            "crear": sum(1 for r in resultado if r["accion"] == "crear"),
            "actualizar": sum(1 for r in resultado if r["accion"] == "actualizar"),
            "sin_cambios": sum(1 for r in resultado if r["accion"] == "sin_cambios"),
            "ignorar": sum(1 for r in resultado if r["accion"] == "ignorar"),
            "con_alertas": sum(1 for r in resultado if r["alertas"]),
        },
        # Estos no se tocan: quitar una fila del archivo no borra el producto.
        "no_presentes": ausentes,
    }


# ---------------------------------------------------------------- aplicar

def aplicar(db: Session, espacio_id, contenido: bytes, usuario_email: str) -> dict:
    """Aplica el archivo. El analisis debe haberse mostrado antes al usuario."""
    filas = leer_filas(contenido)

    existentes = {
        i.nombre.strip().lower(): i
        for i in db.query(Insumo).filter(
            Insumo.espacio_id == espacio_id, Insumo.activo == True
        ).all()
    }
    proveedores = {
        p.nombre.strip().lower(): p
        for p in db.query(Proveedor).filter(Proveedor.espacio_id == espacio_id).all()
    }

    creados, actualizados, detalles = 0, 0, []

    for fila in filas:
        nombre = fila["nombre"]
        if not nombre:
            continue
        clave = nombre.lower()
        insumo = existentes.get(clave)

        if insumo is None:
            insumo = Insumo(
                espacio_id=espacio_id,
                nombre=nombre,
                categoria=fila["categoria"],
                unidad_medida=fila["presentacion"],
                stock_actual=fila["stock_actual"] or Decimal(0),
                stock_minimo=fila["stock_minimo"] or Decimal(0),
                precio_venta=fila["precio_venta"],
                activo=True,
            )
            db.add(insumo)
            db.flush()
            existentes[clave] = insumo
            creados += 1
            detalles.append(f"{nombre}: creado")
        else:
            cambios = []
            if fila["stock_actual"] is not None:
                previo = Decimal(str(insumo.stock_actual or 0))
                if fila["stock_actual"] != previo:
                    cambios.append(f"stock {previo:g} → {fila['stock_actual']:g}")
                    insumo.stock_actual = fila["stock_actual"]
            if fila["stock_minimo"] is not None:
                insumo.stock_minimo = fila["stock_minimo"]
            if fila["precio_venta"] is not None:
                previo = Decimal(str(insumo.precio_venta)) if insumo.precio_venta is not None else None
                if fila["precio_venta"] != previo:
                    cambios.append(f"precio de venta → ${fila['precio_venta']:,.0f}")
                    insumo.precio_venta = fila["precio_venta"]
            if fila["categoria"]:
                insumo.categoria = fila["categoria"]
            if cambios:
                actualizados += 1
                detalles.append(f"{nombre}: " + ", ".join(cambios))

        # El costo entra como un precio nuevo y el anterior se desactiva, de modo
        # que el historico de precios conserve lo que costaba antes.
        if fila["costo"] is not None:
            costo_previo = _costo_actual(insumo)
            if costo_previo is None or fila["costo"] != costo_previo:
                nombre_proveedor = fila["proveedor"] or "Importado desde Excel"
                proveedor = proveedores.get(nombre_proveedor.strip().lower())
                if proveedor is None:
                    proveedor = Proveedor(espacio_id=espacio_id, nombre=nombre_proveedor)
                    db.add(proveedor)
                    db.flush()
                    proveedores[nombre_proveedor.strip().lower()] = proveedor

                for antiguo in insumo.precios:
                    if antiguo.activo and antiguo.proveedor_id == proveedor.id:
                        antiguo.activo = False

                db.add(PrecioInsumo(
                    insumo_id=insumo.id,
                    proveedor_id=proveedor.id,
                    precio_presentacion=fila["costo"],
                    unidades_por_presentacion=1,
                    descripcion_presentacion=fila["presentacion"],
                    activo=True,
                ))

    return {"creados": creados, "actualizados": actualizados, "detalles": detalles}
